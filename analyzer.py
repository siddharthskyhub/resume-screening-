"""Core analysis logic for SkillMatch AI (no Streamlit imports, easy to test)."""

from __future__ import annotations

import io
import re
import unicodedata
from collections import Counter

# --------------------------------------------------------------------------- #
# Skill catalog: category -> skill -> aliases
# Skill names and aliases are matched case-insensitively on word boundaries.
# --------------------------------------------------------------------------- #
SKILL_CATALOG: dict[str, dict[str, list[str]]] = {
    "Languages": {
        "Python": [],
        "Java": [],
        "C++": [],
        "JavaScript": ["ecmascript"],
        "TypeScript": [],
        "SQL": [],
        "HTML": ["html5"],
        "CSS": ["css3"],
    },
    "Web & Backend": {
        "React": ["react.js", "reactjs"],
        "Node.js": ["nodejs", "node js"],
        "Tailwind CSS": ["tailwind"],
        "FastAPI": ["fast api"],
        "Django": [],
        "Flask": [],
        "REST APIs": ["rest api", "restful api", "restful apis"],
    },
    "Data & ML": {
        "Pandas": [],
        "NumPy": [],
        "scikit-learn": ["sklearn", "scikit learn"],
        "PyTorch": [],
        "TensorFlow": [],
        "Machine Learning": ["ml"],
        "Deep Learning": [],
        "NLP": ["natural language processing"],
        "Excel": ["microsoft excel", "ms excel", "advanced excel",
                  "excel spreadsheets", "excel formulas"],
    },
    "GenAI": {
        "LLMs": ["llm", "large language model", "large language models"],
        "RAG": ["retrieval augmented generation",
                "retrieval-augmented generation"],
        "LangChain": ["lang chain"],
        "LangGraph": ["lang graph"],
        "Ollama": [],
        "Hugging Face": ["huggingface"],
        "Prompt Engineering": [],
    },
    "Databases": {
        "MongoDB": ["mongo"],
        "MySQL": [],
        "PostgreSQL": ["postgres"],
        "SQLite": [],
        "Redis": [],
    },
    "Cloud & DevOps": {
        "Docker": [],
        "Kubernetes": ["k8s"],
        "AWS": ["amazon web services"],
        "Azure": [],
        "GCP": ["google cloud"],
        "Linux": [],
        "CI/CD": ["cicd", "ci cd", "continuous integration"],
    },
    "Tools & Concepts": {
        "Git": [],
        "GitHub": [],
        "Streamlit": ["stream lit"],
        "Data Structures": ["dsa", "data structures and algorithms"],
        "OOP": ["object oriented programming", "object-oriented programming"],
    },
}

# Skills whose bare name is too ambiguous (e.g. "excel in communication").
# For these, only the aliases are matched.
NAME_EXCLUDED = {"Excel"}

CATEGORY_OF = {
    skill: category
    for category, skills in SKILL_CATALOG.items()
    for skill in skills
}

SKILL_TIPS = {
    "Python": "Highlight a relevant Python project and what it achieved.",
    "SQL": "Mention the databases you queried and the kind of problems solved.",
    "Git": "Describe your version-control workflow (branches, pull requests).",
    "GitHub": "Link completed repositories with clear README files.",
    "Streamlit": "Show a working demo or screenshots of your app.",
    "REST APIs": "Describe APIs you built or integrations you implemented.",
    "LangChain": "Describe a completed LangChain workflow and its outcome.",
    "LangGraph": "Explain any graph-based agent workflow you have built.",
    "LLMs": "Describe your actual hands-on experience with language models.",
    "RAG": "Explain a RAG project: data source, retrieval method, results.",
    "Docker": "Mention projects you containerized or deployed with Docker.",
    "AWS": "List the specific AWS services you have used.",
    "React": "Describe the components or apps you built with React.",
    "FastAPI": "Mention the APIs you built and how they were consumed.",
    "Pandas": "Highlight data analysis or cleaning work done with Pandas.",
    "NumPy": "Mention numerical computing tasks you handled with NumPy.",
    "Machine Learning": "Describe the models you trained and how you evaluated them.",
}
GENERIC_TIP = (
    "If you genuinely have this skill, add truthful evidence in your "
    "skills, projects, or experience sections."
)
PROOF_FALLBACK = "Add a bullet showing what you built with it and the result."

# --------------------------------------------------------------------------- #
# Text normalization and skill matching
# --------------------------------------------------------------------------- #

def normalize(text: str) -> str:
    """Fix ligatures, collapse whitespace/line breaks, lowercase."""
    text = unicodedata.normalize("NFKC", text)
    return re.sub(r"\s+", " ", text).strip().lower()


def _build_patterns() -> dict[str, re.Pattern]:
    patterns = {}
    for category_skills in SKILL_CATALOG.values():
        for skill, aliases in category_skills.items():
            terms = [] if skill in NAME_EXCLUDED else [skill.lower()]
            terms += [alias.lower() for alias in aliases]
            terms.sort(key=len, reverse=True)  # prefer longest match
            body = "|".join(re.escape(term) for term in terms)
            patterns[skill] = re.compile(rf"(?<!\w)(?:{body})(?!\w)")
    return patterns


PATTERNS = _build_patterns()


def find_skills(text: str) -> dict[str, int]:
    """Return {skill: number_of_mentions} for every catalog skill found."""
    text = normalize(text)
    found = {}
    for skill, pattern in PATTERNS.items():
        count = len(pattern.findall(text))
        if count:
            found[skill] = count
    return found


# --------------------------------------------------------------------------- #
# Job description analysis: required vs. nice-to-have, weighted by frequency
# --------------------------------------------------------------------------- #
PREFERRED_CUES = re.compile(
    r"nice[- ]to[- ]have|preferred|bonus|good to have|a plus|"
    r"optional|desirable|added advantage"
)
REQUIRED_HEADINGS = re.compile(
    r"requirements?|required|must[- ]have|qualifications?|"
    r"responsibilities|what you.ll do|what we.re looking for"
)


def classify_job_skills(job_text: str) -> tuple[dict[str, str], Counter]:
    """Return ({skill: 'required'|'preferred'}, mention_counts)."""
    mode = "required"
    levels: dict[str, str] = {}
    counts: Counter = Counter()

    for chunk in re.split(r"[\r\n]+|(?<=[.!?])\s+", job_text):
        line = normalize(chunk)
        if not line:
            continue

        skills_in_line = find_skills(line)

        # Short lines without skills act as section headings.
        if not skills_in_line and len(line) < 60:
            if PREFERRED_CUES.search(line):
                mode = "preferred"
            elif REQUIRED_HEADINGS.search(line):
                mode = "required"

        level = (
            "preferred"
            if mode == "preferred" or PREFERRED_CUES.search(line)
            else "required"
        )

        for skill, n in skills_in_line.items():
            counts[skill] += n
            if level == "required" or skill not in levels:
                levels[skill] = level

    return levels, counts


def _weight(level: str, count: int) -> float:
    base = 2.0 if level == "required" else 1.0
    return base * (1 + 0.25 * min(count - 1, 3))


def job_profile(job_text: str) -> dict[str, dict]:
    """{skill: {level, count, weight}} for one job description."""
    levels, counts = classify_job_skills(job_text)
    return {
        skill: {"level": level, "count": counts[skill],
                "weight": _weight(level, counts[skill])}
        for skill, level in levels.items()
    }


# --------------------------------------------------------------------------- #
# Resume structure: sections and bullet items
# --------------------------------------------------------------------------- #
HEADINGS = {
    "experience": "experience", "work experience": "experience",
    "professional experience": "experience", "employment history": "experience",
    "work history": "experience", "internship": "experience",
    "internships": "experience",
    "projects": "projects", "project": "projects",
    "personal projects": "projects", "academic projects": "projects",
    "key projects": "projects", "selected projects": "projects",
    "education": "education", "academic background": "education",
    "academics": "education",
    "skills": "skills", "technical skills": "skills", "tech stack": "skills",
    "core competencies": "skills", "skills & tools": "skills",
    "skills and tools": "skills",
    "certifications": "certifications", "certification": "certifications",
    "certificates": "certifications", "courses": "certifications",
    "licenses": "certifications",
    "summary": "summary", "profile": "summary", "objective": "summary",
    "about me": "summary", "professional summary": "summary",
}
BULLET_CHARS = "•●▪◦‣∙·*-–—►▶■□✓➢"


def _match_heading(line: str) -> tuple[str | None, str]:
    """Return (section, remainder) if the line starts a known section."""
    stripped = line.strip()
    core = re.sub(r"[^a-z &/]", "", unicodedata.normalize("NFKC", stripped).lower())
    core = re.sub(r"\s+", " ", core).strip()
    if core in HEADINGS and len(stripped) <= 40:
        return HEADINGS[core], ""

    # Inline form, e.g. "Skills: Python, SQL"
    m = re.match(r"^\W*([A-Za-z][A-Za-z &/]{2,30}?)(?::\s*|\s[-–—]\s+)(.+)$", stripped)
    if m:
        key = re.sub(r"\s+", " ", m.group(1).lower()).strip()
        if key in HEADINGS:
            return HEADINGS[key], m.group(2).strip()
    return None, ""


def split_resume_items(text: str) -> list[tuple[str, str]]:
    """Split a resume into (section, item_text); wrapped bullets are re-joined."""
    items: list[tuple[str, str]] = []
    section = "other"
    current: list | None = None  # [section, text]

    def flush():
        nonlocal current
        if current and current[1].strip():
            items.append((current[0], current[1].strip()))
        current = None

    for raw in unicodedata.normalize("NFKC", text).splitlines():
        line = raw.strip()
        if not line:
            flush()
            continue

        heading, rest = _match_heading(line)
        if heading:
            flush()
            section = heading
            if rest:
                items.append((section, rest))
            continue

        if line[0] in BULLET_CHARS and len(line) > 1:
            flush()
            current = [section, line.lstrip(BULLET_CHARS + " \t")]
        elif current and line[0].islower():
            current[1] += " " + line  # wrapped continuation
        else:
            flush()
            current = [section, line]
    flush()
    return items


def detect_sections(text: str) -> set[str]:
    found = set()
    for line in text.splitlines():
        heading, _ = _match_heading(line)
        if heading:
            found.add(heading)
    return found


# --------------------------------------------------------------------------- #
# Evidence-based matching
# --------------------------------------------------------------------------- #
QUANTIFIED_PATTERN = re.compile(
    r"\d+(?:\.\d+)?\s*%|\$\s?\d|\b\d+(?:\.\d+)?x\b|"
    r"\b\d[\d,.]*[km]?\+?\s+(?:users|requests|records|customers|projects|"
    r"models|ms|hours|stars|downloads)\b"
)
EVIDENCE_SECTIONS = {"experience", "projects", "certifications"}
LEVEL_RANK = {"listed": 0, "demonstrated": 1, "strong": 2}
EVIDENCE_FACTOR = {"strong": 1.0, "demonstrated": 1.0, "listed": 0.5}
LABEL_LINE = re.compile(
    r"^\W*(tech(?:nologies)?(?: stack)?|tools?|stack|skills?)\s*[:\-–—]", re.I
)


# Proving the key skill also proves the underlying one (only upgrades skills
# that already appear on the resume, and only up to "demonstrated").
IMPLIES = {
    "Streamlit": ["Python"], "Flask": ["Python"], "FastAPI": ["Python"],
    "Django": ["Python"], "Pandas": ["Python"], "NumPy": ["Python"],
    "scikit-learn": ["Python"], "PyTorch": ["Python"], "TensorFlow": ["Python"],
    "LangChain": ["Python"], "LangGraph": ["Python"],
    "PostgreSQL": ["SQL"], "MySQL": ["SQL"], "SQLite": ["SQL"],
    "React": ["JavaScript"], "Node.js": ["JavaScript"], "TypeScript": ["JavaScript"],
    "GitHub": ["Git"],
}


def extract_evidence(resume_text: str) -> dict[str, dict]:
    """For each skill on the resume, how well is it backed up?

    levels: 'strong'       used in a project/job bullet WITH a number or result
            'demonstrated' used in a project/job bullet
            'listed'       only appears in a skills list or similar
    """
    evidence: dict[str, dict] = {}

    for section, text in split_resume_items(resume_text):
        found = find_skills(text)
        if not found:
            continue

        words = len(text.split())
        quantified = bool(QUANTIFIED_PATTERN.search(normalize(text)))

        if LABEL_LINE.match(text):
            proves = False
        elif section in EVIDENCE_SECTIONS:
            proves = True
        elif section == "other":
            proves = words >= 7 and len(found) <= 3
        else:
            proves = False

        level = ("strong" if quantified else "demonstrated") if proves else "listed"

        for skill, n in found.items():
            entry = evidence.setdefault(
                skill, {"level": "listed", "mentions": 0, "snippets": []}
            )
            entry["mentions"] += n
            if LEVEL_RANK[level] > LEVEL_RANK[entry["level"]]:
                entry["level"] = level
            if proves:
                entry["snippets"].append(
                    {"section": section, "text": text[:220], "quantified": quantified}
                )

    proven = {k: v for k, v in evidence.items() if v["level"] != "listed"}
    for skill, entry in proven.items():
        for implied in IMPLIES.get(skill, []):
            target = evidence.get(implied)
            if target and target["level"] == "listed" and entry["snippets"]:
                target["level"] = "demonstrated"
                source = entry["snippets"][0]
                target["snippets"].append({
                    "section": source["section"],
                    "text": f"(via {skill}) {source['text']}"[:220],
                    "quantified": False,
                })

    for entry in evidence.values():
        entry["snippets"].sort(key=lambda s: not s["quantified"])
        entry["snippets"] = entry["snippets"][:3]
    return evidence


def score_against(job: dict, resume_found: dict, evidence: dict):
    """Return (keyword_score, evidence_score) as 0-100 ints, or (None, None)."""
    total = sum(info["weight"] for info in job.values())
    if not total:
        return None, None

    keyword = proven = 0.0
    for skill, info in job.items():
        if skill in resume_found:
            keyword += info["weight"]
            level = evidence.get(skill, {}).get("level", "listed")
            proven += info["weight"] * EVIDENCE_FACTOR[level]
    return round(100 * keyword / total), round(100 * proven / total)


# --------------------------------------------------------------------------- #
# Resume quality checks
# --------------------------------------------------------------------------- #

def resume_checks(text: str) -> list[dict]:
    lowered = normalize(text)
    words = len(lowered.split())
    sections = sorted(detect_sections(text) - {"summary"})
    quantified = len(QUANTIFIED_PATTERN.findall(lowered))

    return [
        {
            "label": "Contact email",
            "ok": bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)),
            "detail": "Recruiters need an easy way to reach you.",
        },
        {
            "label": "Profile links",
            "ok": bool(re.search(r"github\.com|linkedin\.com|https?://", lowered)),
            "detail": "Add GitHub, LinkedIn, or a portfolio link.",
        },
        {
            "label": "Standard sections",
            "ok": len(sections) >= 3,
            "detail": (
                "Found: " + ", ".join(sections)
                if sections
                else "Use clear headings like Experience, Projects, Education, Skills."
            ),
        },
        {
            "label": "Quantified achievements",
            "ok": quantified >= 2,
            "detail": f"{quantified} found. Numbers (%, users, time saved) make impact concrete.",
        },
        {
            "label": "Length",
            "ok": 200 <= words <= 1000,
            "detail": f"{words} words. Aim for roughly 300-800 for most resumes.",
        },
    ]


# --------------------------------------------------------------------------- #
# Main analysis (one resume vs one job)
# --------------------------------------------------------------------------- #

def analyze_resume(resume_text: str, job_text: str) -> dict:
    resume_found = find_skills(resume_text)
    evidence = extract_evidence(resume_text)
    job = job_profile(job_text)

    ordered = sorted(job, key=lambda s: (-job[s]["weight"], s))
    matched = [s for s in ordered if s in resume_found]
    missing = [s for s in ordered if s not in resume_found]
    extra = sorted(s for s in resume_found if s not in job)
    unproven = [
        s for s in matched if evidence.get(s, {}).get("level", "listed") == "listed"
    ]

    score, evidence_score = score_against(job, resume_found, evidence)

    by_category = {}
    for category in SKILL_CATALOG:
        in_job = [s for s in job if CATEGORY_OF[s] == category]
        if in_job:
            hits = sum(1 for s in in_job if s in resume_found)
            by_category[category] = (hits, len(in_job))

    suggestions = [(skill, SKILL_TIPS.get(skill, GENERIC_TIP)) for skill in missing]
    proof_suggestions = [
        (skill,
         "Listed, but no project or job bullet shows it. "
         + SKILL_TIPS.get(skill, PROOF_FALLBACK))
        for skill in unproven
    ]

    return {
        "score": score,
        "evidence_score": evidence_score,
        "resume_skills": resume_found,
        "evidence": evidence,
        "job": job,
        "matched": matched,
        "missing": missing,
        "unproven": unproven,
        "extra": extra,
        "by_category": by_category,
        "suggestions": suggestions,
        "proof_suggestions": proof_suggestions,
        "checks": resume_checks(resume_text),
        "resume_text": resume_text,
    }


def build_report(result: dict) -> str:
    def pct(v):
        return "n/a" if v is None else f"{v}%"

    lines = [
        "SKILLMATCH AI - RESUME ANALYSIS",
        "=" * 32,
        "",
        f"Keyword match: {pct(result['score'])}",
        f"Evidence-backed match: {pct(result['evidence_score'])}",
        f"Resume skills detected: {len(result['resume_skills'])}",
        f"Matching skills: {', '.join(result['matched']) or 'None'}",
        f"Listed but unproven: {', '.join(result['unproven']) or 'None'}",
        f"Skills to review: {', '.join(result['missing']) or 'None'}",
        f"Extra strengths: {', '.join(result['extra']) or 'None'}",
        "",
        "Suggestions:",
    ]
    lines += [f"- {s}: {t}" for s, t in result["suggestions"] + result["proof_suggestions"]] or ["- None"]
    lines += ["", "Resume checks:"]
    lines += [
        f"- [{'OK' if c['ok'] else 'REVIEW'}] {c['label']}: {c['detail']}"
        for c in result["checks"]
    ]
    lines += ["", "Note: this is keyword matching, not a prediction of hiring success."]
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# File extraction
# --------------------------------------------------------------------------- #

class ResumeError(Exception):
    """Raised with a user-friendly message when a resume can't be read."""


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ResumeError("This PDF is password-protected. Please upload an unlocked copy.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ResumeError:
        raise
    except Exception as exc:
        raise ResumeError("That PDF could not be read. It may be corrupted.") from exc


def _extract_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    except Exception as exc:
        raise ResumeError("That DOCX file could not be read.") from exc


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = _extract_pdf(data)
    elif name.endswith(".docx"):
        text = _extract_docx(data)
    else:
        raise ResumeError("Unsupported file type. Please upload a PDF or DOCX.")

    if not text.strip():
        raise ResumeError(
            "No readable text was found. The file may contain scanned images."
        )
    return text
